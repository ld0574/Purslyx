"""资料解析、格式识别和异常输入的单元测试。"""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from pypdf import PdfWriter
from reportlab.pdfgen import canvas

from server.app import parsing
from server.app.errors import DomainError
from server.app.parsing import (
    detect_source_type,
    extract_file_text,
    normalize_text,
    parse_job_text,
    parse_resume_text,
    parse_salary_text,
    sha256_bytes,
)


def test_normalize_and_hash_are_stable() -> None:
    assert normalize_text("  第一行  \r\n第二行\t \r\n") == "第一行\n第二行"
    assert sha256_bytes(b"purslyx") == sha256_bytes(b"purslyx")
    assert sha256_bytes(b"purslyx") != sha256_bytes(b"Purslyx")


@pytest.mark.parametrize(
    ("filename", "content_type", "expected"),
    [
        ("resume.TXT", None, "text"),
        ("resume.pdf", "application/octet-stream", "pdf"),
        ("resume.doc", None, "doc"),
        ("resume.docx", None, "docx"),
        (None, "text/plain", "text"),
        (None, "application/pdf", "pdf"),
        (
            None,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            "docx",
        ),
    ],
)
def test_detect_source_type(filename: str | None, content_type: str | None, expected: str) -> None:
    assert detect_source_type(filename, content_type) == expected


def test_detect_source_type_rejects_images() -> None:
    with pytest.raises(DomainError) as error:
        detect_source_type("scan.png", "image/png")
    assert error.value.code == "DOCUMENT_FORMAT_UNSUPPORTED"
    assert error.value.status_code == 415


def test_extract_text_and_empty_file_contract(tmp_path: Path) -> None:
    valid = tmp_path / "resume.txt"
    valid.write_text("姓名：测试用户\n后端开发经历", encoding="utf-8")
    assert extract_file_text(valid, "text") == "姓名：测试用户\n后端开发经历"

    empty = tmp_path / "empty.txt"
    empty.write_bytes(b"  \n")
    with pytest.raises(DomainError) as error:
        extract_file_text(empty, "text")
    assert error.value.code == "DOCUMENT_CONTENT_UNREADABLE"
    assert error.value.action == "paste_text"


def test_extract_rejects_file_and_text_over_limit(
    monkeypatch: pytest.MonkeyPatch,
    tmp_path: Path,
) -> None:
    file_path = tmp_path / "large.txt"
    file_path.write_bytes(b"12345")
    monkeypatch.setattr(parsing, "MAX_FILE_BYTES", 4)
    with pytest.raises(DomainError) as file_error:
        extract_file_text(file_path, "text")
    assert file_error.value.code == "DOCUMENT_TOO_LARGE"

    monkeypatch.setattr(parsing, "MAX_FILE_BYTES", 100)
    monkeypatch.setattr(parsing, "MAX_TEXT_CHARS", 4)
    with pytest.raises(DomainError) as text_error:
        extract_file_text(file_path, "text")
    assert text_error.value.code == "DOCUMENT_TOO_LARGE"


def test_extract_docx_includes_paragraphs_and_tables(tmp_path: Path) -> None:
    file_path = tmp_path / "resume.docx"
    document = Document()
    document.add_paragraph("Purslyx candidate")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Skill"
    table.cell(0, 1).text = "Python"
    document.save(file_path)

    text = extract_file_text(file_path, "docx")
    assert "Purslyx candidate" in text
    assert "Skill | Python" in text


def test_extract_docx_includes_text_boxes(tmp_path: Path) -> None:
    file_path = tmp_path / "template.docx"
    document = Document()
    paragraph = document.add_paragraph("姓名\t")
    textbox = OxmlElement("w:txbxContent")
    box_paragraph = OxmlElement("w:p")
    box_run = OxmlElement("w:r")
    box_text = OxmlElement("w:t")
    box_text.text = "负责 Python 项目"
    box_run.append(box_text)
    box_paragraph.append(box_run)
    textbox.append(box_paragraph)
    paragraph._p.append(textbox)
    document.sections[0].header.paragraphs[0].text = "求职意向：研发工程师"
    document.save(file_path)

    assert "姓名\t负责 Python 项目" in extract_file_text(file_path, "docx")
    assert "求职意向：研发工程师" in extract_file_text(file_path, "docx")


def test_extract_pdf_and_reject_encrypted_pdf(tmp_path: Path) -> None:
    readable = tmp_path / "readable.pdf"
    output = canvas.Canvas(str(readable))
    output.drawString(72, 760, "Purslyx PDF resume")
    output.save()
    assert "Purslyx PDF resume" in extract_file_text(readable, "pdf")

    encrypted = tmp_path / "encrypted.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt("secret")
    with encrypted.open("wb") as handle:
        writer.write(handle)
    with pytest.raises(DomainError) as error:
        extract_file_text(encrypted, "pdf")
    assert error.value.code == "DOCUMENT_CONTENT_UNREADABLE"
    assert error.value.action == "paste_text"


def test_extract_pdf_uses_second_reader_when_first_has_no_text(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    readable = tmp_path / "readable.pdf"
    output = canvas.Canvas(str(readable))
    output.drawString(72, 760, "PDF fallback resume")
    output.save()

    class EmptyReader:
        def __init__(self, _: str) -> None:
            self.pages = [type("Page", (), {"extract_text": lambda self: "\x01"})()]

    monkeypatch.setattr("pypdf.PdfReader", EmptyReader)
    assert "PDF fallback resume" in extract_file_text(readable, "pdf")

    def broken_reader(_: str):
        raise ValueError("unrecognized font mapping")

    monkeypatch.setattr("pypdf.PdfReader", broken_reader)
    assert "PDF fallback resume" in extract_file_text(readable, "pdf")


def test_legacy_doc_uses_utf8_reader(monkeypatch: pytest.MonkeyPatch, tmp_path: Path) -> None:
    file_path = tmp_path / "resume.doc"
    file_path.write_bytes(b"old Word file")
    calls: list[list[str]] = []
    monkeypatch.setattr(parsing.shutil, "which", lambda name: "/usr/bin/antiword" if name == "antiword" else None)

    def fake_run(args: list[str], **_: object):
        calls.append(args)
        return type("Result", (), {"returncode": 0, "stdout": "工作经历：\n负责中文简历"})()

    monkeypatch.setattr(parsing.subprocess, "run", fake_run)
    assert "负责中文简历" in extract_file_text(file_path, "doc")
    assert calls == [["/usr/bin/antiword", "-m", "UTF-8.txt", str(file_path)]]


def test_resume_parser_preserves_only_supplied_segments() -> None:
    result = parse_resume_text("张三\n工作经历：\n负责 API 设计\n项目经历\nPurslyx 项目")
    segments = [
        segment["text"]
        for section in result["sections"]
        for segment in section["segments"]
    ]
    assert segments == ["张三", "负责 API 设计", "Purslyx 项目"]
    assert all(segment["source"] == "user_confirmed" for section in result["sections"] for segment in section["segments"])


def test_resume_parser_keeps_repeated_headings_unique() -> None:
    result = parse_resume_text("Experience:\n负责项目 A\nExperience:\n负责项目 B")
    assert [section["section_key"] for section in result["sections"]] == ["experience", "experience-2"]
    assert [section["segments"][0]["text"] for section in result["sections"]] == ["负责项目 A", "负责项目 B"]


def test_job_parser_keeps_missing_fields_unknown() -> None:
    result = parse_job_text(
        "高级后端开发工程师\n公司：示例科技\n地点：杭州／上海\n薪资：20K-30K/月\n3-5年，本科\n现场办公\n1、熟悉 Python\n负责 API 开发"
    )["job_fields"]
    assert result["title"] == "高级后端开发工程师"
    assert result["company_name"] == "示例科技"
    assert result["locations"] == ["杭州", "上海"]
    assert result["work_mode"] == "onsite"
    assert result["experience_text"] == "3-5年，本科"
    assert result["education_text"] == "3-5年，本科"
    assert result["salary"]["min"] == "20000.0"
    assert result["salary"]["max"] == "30000.0"
    assert result["salary"]["currency"] is None
    assert result["category"] == "engineering"


@pytest.mark.parametrize(
    ("value", "status"),
    [
        (None, None),
        ("薪资面议", "negotiable"),
        ("按经验确定", "unknown"),
        ("15-25K/月", "specified"),
    ],
)
def test_salary_parser_does_not_invent_basis(value: str | None, status: str | None) -> None:
    result = parse_salary_text(value)
    if status is None:
        assert result is None
    else:
        assert result is not None
        assert result["status"] == status
        assert result["currency"] is None
        assert result["tax_basis"] is None
